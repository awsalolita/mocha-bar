#!/usr/bin/env python3
"""Generate a WorldSkills CIS-format Marking Scheme for the FactoryOps
(WSC2026_TP53) Cloud Computing test project.

The layout mirrors the standard "CIS Marking Scheme Import" sheet:
  - Skill title
  - WorldSkills Standards Specification (WSSS) section table with
    WSSS Marks / Aspect Marks / Variation
  - Criteria table (A..D) with per-criterion marks
  - Per-criterion aspect blocks (Measurement M = Yes/No, Judgement J = 0..3)

Aspect max marks drive everything: per-criterion totals (column N) and the
per-WSSS-section aspect sums are computed so Variation is always 0 and the
grand total is the sum of all aspects.

Run:  python generate_marking.py
Out:  WSC2026_TP53_MarkingScheme.xlsx
"""

import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# --- WSSS sections (standard Cloud Computing specification) ----------------
WSSS = {
    1: "Work organization and management",
    2: "Communication and interpersonal skills",
    3: "Problem solving, innovation, and creativity",
    4: "Cybersecurity",
    5: "Reliability, scalability, and elasticity",
    6: "Performance and optimization",
    7: "Operational considerations",
    8: "Sustainability",
}


def M(desc, extra, wsss, mark):
    """Measurement aspect (objective Yes/No)."""
    return {"type": "M", "desc": desc, "extra": extra, "wsss": wsss, "mark": mark}


def J(desc, wsss, mark, scores):
    """Judgement aspect scored 0..3 with per-score descriptions."""
    return {"type": "J", "desc": desc, "wsss": wsss, "mark": mark, "scores": scores}


# --- Marking content --------------------------------------------------------
# Each criterion: (id, name, day_of_marking, [ (subid, subname, [aspects]) ])
CRITERIA = [
    ("A", "FactoryOps - Network Foundation", 1, [
        ("A1", "VPC", [
            M("Are three VPCs created (VPC-EDGE, VPC-APP, VPC-DATA)",
              "Three separate VPCs exist as per the architecture.", 5, 1),
            M("Are subnets deployed across multiple Availability Zones",
              "Removes single points of failure (multi-AZ).", 5, 1),
            M("Are non-overlapping CIDR ranges used for the three VPCs",
              "Enables clean routing over Transit Gateway.", 1, 1),
            J("Overall VPC/subnet network design quality", 3, 3, [
                "No coherent VPC design.",
                "VPCs exist but subnetting is flat or inconsistent.",
                "VPCs with public/private subnet tiers in one or two VPCs.",
                "Clean 3-VPC design with consistent public/private tiers and AZ spread.",
            ]),
        ]),
        ("A2", "Transit Gateway", [
            M("Is a Transit Gateway created", "A TGW exists in us-east-1.", 5, 1),
            M("Are all three VPCs attached to the Transit Gateway",
              "EDGE, APP and DATA attachments exist.", 5, 1),
            M("Are VPC and TGW route tables configured for required paths",
              "Routes allow the intended cross-VPC communication.", 5, 1),
            J("Transit Gateway routing correctness", 5, 2, [
                "No working cross-VPC routing.",
                "Partial routing; some required paths fail.",
                "All required paths work but routing is overly permissive.",
                "All required paths work with least-privilege TGW route tables.",
            ]),
        ]),
        ("A3", "Segmentation", [
            J("Direct communication between VPC-EDGE and VPC-DATA is blocked", 4, 2, [
                "EDGE can reach DATA directly (segmentation absent).",
                "Partial block; some EDGE->DATA traffic still allowed.",
                "EDGE->DATA blocked but via ad-hoc/incorrect mechanism.",
                "EDGE<->DATA fully blocked via dedicated TGW route tables.",
            ]),
            M("Is VPC-EDGE to VPC-APP communication allowed",
              "ui-service (EDGE) can reach core-service (APP).", 4, 1),
            M("Is VPC-APP to VPC-DATA communication allowed",
              "core/worker path (APP<->DATA) works.", 4, 1),
            M("Is cross-boundary traffic auditable (e.g. VPC Flow Logs)",
              "Communication crossing VPC boundaries is logged/auditable.", 4, 1),
        ]),
        ("A4", "Cloud Map / Private DNS", [
            M("Is a Cloud Map private DNS namespace created (svc.internal)",
              "Private namespace exists for service discovery.", 7, 1),
            M("Does core.svc.internal resolve", "FQDN resolves to core-service.", 7, 0.5),
            M("Does worker.svc.internal resolve", "FQDN resolves to worker-service.", 7, 0.5),
            M("Are services discoverable across VPCs via FQDN",
              "Cross-VPC discovery works over the namespace.", 5, 1),
        ]),
        ("A5", "Private access to AWS services", [
            M("Are VPC endpoints created for required AWS services",
              "Endpoints for SQS, DynamoDB, ECR, etc.", 4, 1),
            M("Is public internet exposure minimized for private workloads",
              "Private subnets avoid unnecessary IGW/NAT exposure.", 4, 1),
            J("Private connectivity design to AWS services", 4, 1, [
                "Workloads reach AWS services over the public internet.",
                "Some services via endpoints, others public.",
                "Most required services via VPC endpoints.",
                "All required services reached privately via VPC endpoints.",
            ]),
        ]),
    ]),

    ("B", "FactoryOps - Compute & Services", 1, [
        ("B1", "EC2 / ui-service", [
            M("Is an EC2 instance deployed in VPC-EDGE", "ui-service host in EDGE.", 7, 1),
            M("Is ui-service listening on port 8080", "Correct listening port.", 7, 0.5),
            M("Does GET /health return healthy", "Health check endpoint works.", 7, 0.5),
            M("Is the web UI for event visualization accessible",
              "UI renders event data.", 7, 1),
            M("Does POST /api/events forward events to core-service",
              "Entry endpoint forwards for processing.", 5, 1),
            J("EC2 host hardening (IMDSv2, security groups)", 4, 1, [
                "IMDSv1 open and permissive security groups.",
                "Some hardening applied.",
                "IMDSv2 enforced or tight SGs, not both.",
                "IMDSv2 enforced and least-privilege security groups.",
            ]),
        ]),
        ("B2", "ECS Clusters", [
            M("Does app-cluster exist in VPC-APP", "ECS cluster in APP.", 7, 1),
            M("Does worker-cluster exist in VPC-DATA", "ECS cluster in DATA.", 7, 1),
            M("Do ECS tasks send logs to stdout/stderr (awslogs)",
              "Container logs available for troubleshooting.", 7, 1),
        ]),
        ("B3", "core-service", [
            M("Is core-service running in app-cluster", "Service deployed in APP.", 7, 1),
            M("Is core-service registered as core.svc.internal",
              "Registered in Cloud Map.", 7, 0.5),
            M("Are Type-A events routed to worker-service (synchronous)",
              "Sync call to worker.svc.internal.", 5, 1),
            M("Are Type-B events routed to SQS (asynchronous)",
              "Enqueued to events-b-queue.", 5, 0.5),
            J("Event routing correctness and robustness", 3, 1, [
                "Routing does not work.",
                "One event type routed correctly.",
                "Both types routed but error handling weak.",
                "Both types routed with sound validation and error handling.",
            ]),
        ]),
        ("B4", "worker-service", [
            M("Is worker-service running in worker-cluster", "Service deployed in DATA.", 7, 1),
            M("Is worker-service registered as worker.svc.internal",
              "Registered in Cloud Map.", 7, 0.5),
            M("Is worker-service reachable from core over the required path",
              "APP->DATA path validated.", 5, 0.5),
            J("worker-service deployment quality", 7, 1, [
                "Not deployed / not reachable.",
                "Deployed but unstable.",
                "Deployed and reachable.",
                "Deployed, reachable and healthy with health checks.",
            ]),
        ]),
        ("B5", "Single entry endpoint", [
            M("Is a single entry endpoint provided (e.g. ALB)",
              "One entry endpoint rather than many.", 7, 2),
            M("Are the required communication paths validated",
              "End-to-end paths verified.", 5, 1),
            J("Entry endpoint design", 7, 1, [
                "No single entry endpoint.",
                "Multiple endpoints exposed.",
                "Single endpoint but not fronted by a load balancer.",
                "Single load-balanced entry endpoint.",
            ]),
        ]),
        ("B6", "Auto scaling", [
            M("Is ECS service auto scaling configured", "Scaling policy exists.", 5, 1),
            J("Workload capacity adjusts to demand", 5, 3, [
                "No scaling policy exists.",
                "Scaling policy on a single service.",
                "Scaling policies on core and worker services.",
                "Scaling policies exist and worked according to the policy.",
            ]),
            M("Are min/max limits sensible to avoid over-scaling",
              "Cost-aware scaling bounds.", 6, 1),
        ]),
    ]),

    ("C", "FactoryOps - Event Pipeline & Data", 2, [
        ("C1", "SQS", [
            M("Is the events-b-queue created", "Main queue exists.", 5, 1),
            M("Is the events-b-dlq created", "Dead-letter queue exists.", 5, 1),
            M("Is a redrive (DLQ) policy configured on the queue",
              "Failed messages move to DLQ.", 5, 1),
            J("Queue reliability configuration", 5, 2, [
                "Queue not usable.",
                "Queue exists without DLQ wiring.",
                "Queue + DLQ with default settings.",
                "Queue + DLQ with appropriate visibility timeout/retention/redrive.",
            ]),
        ]),
        ("C2", "Lambda event-normalizer", [
            M("Is the Lambda named event-normalizer and SQS-triggered",
              "Consumes from events-b-queue.", 6, 1),
            M("Does it validate orderId starts with ORD-",
              "Type-B order id rule.", 3, 0.5),
            M("Does it validate operator starts with op- or operator-",
              "Type-B operator rule.", 3, 0.5),
            M("Does it validate operation is PACK or SHIP",
              "Type-B operation rule.", 3, 0.5),
            M("Does it normalize fields (timestamp handling, trimming)",
              "Normalization applied before persistence.", 3, 1),
            M("Does it write to DynamoDB idempotently",
              "Duplicate/out-of-order events do not double-write.", 4, 1.5),
            M("Does it raise an exception on failure (retry/DLQ applies)",
              "Errors trigger SQS retry then DLQ.", 4, 1),
            J("Validation and normalization quality", 3, 2, [
                "No validation/normalization.",
                "Partial validation only.",
                "All rules validated, normalization incomplete.",
                "All rules validated and fields fully normalized.",
            ]),
        ]),
        ("C3", "Aurora + RDS Proxy", [
            M("Is Aurora PostgreSQL provisioned (not serverless)",
              "Provisioned cluster as required.", 5, 1),
            M("Does the events_orders table match the required schema",
              "event_id PK, order_id, amount, ts, created_at.", 3, 1),
            M("Is RDS Proxy configured", "Proxy fronts the Aurora cluster.", 5, 1),
            M("Does worker-service connect via RDS Proxy (not directly)",
              "Connection goes through the proxy endpoint.", 4, 1),
            J("Aurora high availability and backup configuration", 8, 2, [
                "Single instance, no backups.",
                "Backups enabled only.",
                "Multi-AZ or backups with retention, not both.",
                "Multi-AZ with backups and adequate retention.",
            ]),
        ]),
        ("C4", "DynamoDB", [
            M("Is the events_operations table created", "Table exists.", 5, 1),
            M("Is the partition key event_id (String)", "Correct PK.", 3, 1),
            M("Is the sort key ts (String, ISO8601)", "Correct SK.", 3, 1),
            M("Is backup / PITR enabled on the table",
              "Recovery options for sustainability.", 8, 1),
        ]),
        ("C5", "Type-A pipeline (end-to-end)", [
            M("Is a Type-A event persisted to Aurora via the pipeline",
              "ui -> core -> worker -> RDS Proxy -> Aurora.", 5, 2),
            M("Are duplicate Type-A events handled idempotently",
              "No duplicate rows in events_orders.", 5, 1),
            J("Reliable processing of out-of-order / duplicated Type-A events", 5, 1, [
                "Duplicates create duplicate rows / errors.",
                "Some duplicates handled.",
                "Duplicates ignored but out-of-order mishandled.",
                "Duplicates and out-of-order events handled correctly.",
            ]),
        ]),
        ("C6", "Type-B pipeline (end-to-end)", [
            M("Does a valid Type-B event reach DynamoDB",
              "ui -> core -> SQS -> Lambda -> DynamoDB.", 5, 2),
            M("Are invalid Type-B events routed to the DLQ",
              "Validation failures land in events-b-dlq.", 4, 1),
            M("Are duplicate Type-B events handled idempotently",
              "No duplicate items in events_operations.", 5, 1),
        ]),
        ("C7", "Read-only API", [
            M("Does GET /api/events return application/json", "Correct response format.", 6, 1),
            M("Does the response consolidate Type-A and Type-B datasets",
              "Both datasets in one view.", 6, 1),
            J("Read-only API correctness and completeness", 6, 1, [
                "Endpoint missing or broken.",
                "Returns one dataset only.",
                "Returns both but shape is inconsistent.",
                "Returns a clean consolidated view of both datasets.",
            ]),
        ]),
    ]),

    ("D", "FactoryOps - Security & Operations", 2, [
        ("D1", "Data protection", [
            M("Is encryption at rest enabled (KMS) for the datastores",
              "Aurora and DynamoDB encrypted.", 4, 1),
            M("Is data protected in transit (TLS / DB SSL)",
              "HTTPS and encrypted DB connections.", 4, 1),
            M("Are database credentials stored in Secrets Manager",
              "No plaintext credentials.", 4, 1),
            J("Data protection across the lifecycle", 4, 2, [
                "Data unprotected.",
                "Partial protection (rest or transit).",
                "Rest and transit protected with AWS-managed keys.",
                "Rest and transit protected with CMKs and managed secrets.",
            ]),
        ]),
        ("D2", "Account security & governance", [
            M("Is CloudTrail enabled", "Auditable API activity.", 4, 1),
            M("Is AWS Config or GuardDuty enabled", "Account-wide detective controls.", 4, 1),
            M("Do IAM roles follow least privilege", "Scoped task/execution roles.", 4, 1),
            J("Overall account security posture", 4, 2, [
                "No account-wide controls.",
                "One control enabled.",
                "Several controls with broad IAM.",
                "Comprehensive controls with least-privilege IAM.",
            ]),
        ]),
        ("D3", "Observability", [
            M("Do CloudWatch log groups exist for the services",
              "Logs collected for key components.", 7, 1),
            M("Are key performance metrics collected",
              "CPU/latency/queue metrics captured.", 6, 1),
            J("Monitoring dashboards completeness", 7, 2, [
                "No dashboard exists.",
                "A dashboard exists with few metrics.",
                "Dashboard shows most key metrics.",
                "Dashboard shows all key metrics across components.",
            ]),
        ]),
        ("D4", "Traceability", [
            M("Is request tracing implemented (e.g. X-Ray / request id)",
              "Requests can be traced across services.", 7, 1),
            J("End-to-end request traceability", 7, 2, [
                "No traceability.",
                "Partial tracing on some services.",
                "Tracing across most of the path.",
                "Full end-to-end traceability (ui -> core -> worker/SQS).",
            ]),
        ]),
        ("D5", "Cost & Well-Architected", [
            M("Are resources right-sized without excessive scaling",
              "Cost monitored and minimized.", 6, 1),
            J("Alignment with the AWS Well-Architected Framework", 6, 2, [
                "No consideration of Well-Architected pillars.",
                "One or two pillars addressed.",
                "Most pillars addressed.",
                "Solution optimized across all pillars.",
            ]),
        ]),
    ]),
]

SKILL_TITLE = "53 Cloud Computing"
PROJECT_TITLE = "WSC2026_TP53 - FactoryOps Event Management Platform"

# --- Styling ---------------------------------------------------------------
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
GRP_FILL = PatternFill("solid", fgColor="FCE4D6")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COL_HEADERS = {
    "A": "Sub\nCriterion\nID",
    "B": "Sub Criterion\nName or Description",
    "C": "Day of Marking",
    "D": "Aspect\nType\nM = Meas\nJ = Judg",
    "E": "Aspect - Description",
    "F": "Judg Score",
    "G": "Extra Aspect Description (Meas or Judg)\nOR\nJudgement Score Description (Judg only)",
    "H": "Requirement\n(Measurement Only)",
    "I": "WSSS Section",
    "J": "Calculation Row \n(Export only)",
    "K": "Max\nMark",
}


def aspect_max(a):
    return a["mark"]


def criterion_total(crit):
    _, _, _, subs = crit
    return sum(aspect_max(a) for _, _, aspects in subs for a in aspects)


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "CIS Marking Scheme Import"

    # per-section aspect totals for the WSSS table
    section_marks = defaultdict(float)
    for crit in CRITERIA:
        for _, _, aspects in crit[3]:
            for a in aspects:
                section_marks[a["wsss"]] += a["mark"]

    r = 1
    ws.cell(r, 1, SKILL_TITLE).font = TITLE_FONT
    r += 1
    ws.cell(r, 1, PROJECT_TITLE).font = Font(italic=True, color="808080")
    r += 2

    ws.cell(r, 1, "WorldSkills Standards Specification").font = BOLD
    r += 1
    # WSSS table header
    ws.cell(r, 1, "Section").font = BOLD
    ws.cell(r, 9, "WSSS Marks").font = BOLD
    ws.cell(r, 10, "Aspect Marks").font = BOLD
    ws.cell(r, 11, "Variation").font = BOLD
    r += 1
    grand = 0.0
    for sec in range(1, 9):
        marks = round(section_marks.get(sec, 0), 2)
        grand += marks
        ws.cell(r, 1, sec)
        ws.cell(r, 2, WSSS[sec])
        ws.cell(r, 9, marks)       # WSSS target = computed aspect marks
        ws.cell(r, 10, marks)      # aspect marks
        ws.cell(r, 11, 0)          # variation
        r += 1
    ws.cell(r, 9, "Total").font = BOLD
    ws.cell(r, 10, round(grand, 2)).font = BOLD
    ws.cell(r, 11, "Total Variation")
    ws.cell(r, 12, 0)
    r += 2

    # Criteria summary table
    ws.cell(r, 1, "Criteria").font = BOLD
    r += 1
    ws.cell(r, 1, "ID").font = BOLD
    ws.cell(r, 2, "Name").font = BOLD
    ws.cell(r, 11, "Mark").font = BOLD
    r += 1
    for crit in CRITERIA:
        cid, cname, _, _ = crit
        ws.cell(r, 1, cid)
        ws.cell(r, 2, cname)
        ws.cell(r, 11, round(criterion_total(crit), 2))
        r += 1
    r += 1

    # Per-criterion aspect blocks
    for crit in CRITERIA:
        cid, cname, day, subs = crit
        total = round(criterion_total(crit), 2)

        # block header row
        for col, text in COL_HEADERS.items():
            c = ws.cell(r, ord(col) - 64, text)
            c.font = WHITE_BOLD
            c.fill = HDR_FILL
            c.alignment = CENTER
            c.border = BORDER
        ws.cell(r, 12, f"Criterion {cid}").font = WHITE_BOLD
        ws.cell(r, 12).fill = HDR_FILL
        ws.cell(r, 12).alignment = CENTER
        ws.cell(r, 13, "Total\nMark").font = WHITE_BOLD
        ws.cell(r, 13).fill = HDR_FILL
        ws.cell(r, 13).alignment = CENTER
        ws.cell(r, 14, total).font = WHITE_BOLD
        ws.cell(r, 14).fill = HDR_FILL
        ws.cell(r, 14).alignment = CENTER
        r += 1

        for subid, subname, aspects in subs:
            # sub-criterion group row
            gc = ws.cell(r, 1, subid)
            gc.font = BOLD
            ws.cell(r, 2, subname).font = BOLD
            ws.cell(r, 3, day)
            for col in range(1, 12):
                ws.cell(r, col).fill = GRP_FILL
            r += 1

            for a in aspects:
                if a["type"] == "M":
                    ws.cell(r, 4, "M").alignment = CENTER
                    ws.cell(r, 5, a["desc"]).alignment = LEFT
                    ws.cell(r, 7, a["extra"]).alignment = LEFT
                    ws.cell(r, 8, "Yes/No").alignment = CENTER
                    ws.cell(r, 9, a["wsss"]).alignment = CENTER
                    ws.cell(r, 11, a["mark"]).alignment = CENTER
                    r += 1
                else:  # Judgement
                    ws.cell(r, 4, "J").alignment = CENTER
                    ws.cell(r, 5, a["desc"]).alignment = LEFT
                    ws.cell(r, 9, a["wsss"]).alignment = CENTER
                    ws.cell(r, 11, a["mark"]).alignment = CENTER
                    r += 1
                    for score, text in enumerate(a["scores"]):
                        ws.cell(r, 6, score).alignment = CENTER
                        ws.cell(r, 7, text).alignment = LEFT
                        r += 1
            r += 1  # blank row after sub-criterion

    # footer total
    ws.cell(r, 12, "Competition").font = BOLD
    ws.cell(r, 13, "Total\nMark").font = BOLD
    ws.cell(r, 13).alignment = CENTER
    ws.cell(r, 14, round(grand, 2)).font = BOLD

    # column widths
    widths = {"A": 10, "B": 30, "C": 8, "D": 10, "E": 46, "F": 6,
              "G": 52, "H": 12, "I": 8, "J": 10, "K": 7, "L": 14, "M": 9, "N": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "WSC2026_TP53_MarkingScheme.xlsx")
    wb.save(out)
    print(f"Saved: {out}")
    print(f"Grand total: {round(grand, 2)}")
    for crit in CRITERIA:
        print(f"  Criterion {crit[0]} ({crit[1]}): {round(criterion_total(crit), 2)}")


if __name__ == "__main__":
    build()
